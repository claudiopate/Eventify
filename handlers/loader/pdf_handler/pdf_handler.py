from telegram import Update
from telegram.ext import ContextTypes, ConversationHandler
import os
from img2table.document import PDF
from img2table.tables.objects.table import Table
from io import BytesIO
from img2table.ocr import TesseractOCR
from IPython.display import display_html
from openpyxl import load_workbook


tesseract_ocr = TesseractOCR(n_threads=1, lang="ita")

#Status of waiting
WAITING_FOR_PDF = 1

# Funzione per gestire il caricamento del PDF
async def handle_pdf(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    file = update.message.document

    # Verifica se il file è un PDF
    if file.mime_type == 'application/pdf':
        # Scarica il file PDF
        file_extension = ".pdf"
        file_id = file.file_id
        new_file = await context.bot.get_file(file_id)
        file_path = os.path.join("download", f"{file_id}{file_extension}")
        await new_file.download_to_drive(file_path)

        await update.message.reply_text("PDF ricevuto! Ora lo processerò.")
        
        # Qui puoi aggiungere il codice per processare il PDF
        
        extracted_tables = extract_table_from_pdf(file_path)
        
        for page, tables in extracted_tables.items():
            for idx, table in enumerate(tables):
                display_html(table.html_repr(title=f"Page {page + 1} - Extracted table n°{idx + 1}"), raw=True)
        
        return ConversationHandler.END  # Fine della conversazione
    else:
        await update.message.reply_text("Il file inviato non è un PDF. Per favore, invia un file PDF.")
        return WAITING_FOR_PDF  # Resta nello stato di attesa del PDF
    
def extract_table_from_pdf(pdf_path: str):
    
    pdf = PDF(pdf_path, 
          pages=[0],
          detect_rotation=False,
          pdf_text_extraction=True)
    
    with open(pdf_path, 'rb') as f:
        pdf_bytes = f.read()
        
    pdf_from_bytes = PDF(src=pdf_bytes)
    
    pdf_from_file_like = PDF(src=BytesIO(pdf_bytes))
    
    # Extract tables
    extracted_tables = pdf.extract_tables(ocr=tesseract_ocr,
                                      implicit_rows=True,
                                      borderless_tables=False,
                                      min_confidence=50)
    
    pdf.to_xlsx('data/tables.xlsx',
            ocr=tesseract_ocr,
            implicit_rows=True,
            borderless_tables=False,
            min_confidence=50)
    
    for ws in load_workbook("data/tables.xlsx"):
        print(f"Worksheet {ws.title} : {len(tuple(ws.rows))} rows, {len(tuple(ws.rows)[0])} columns")
